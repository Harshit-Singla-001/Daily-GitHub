import os

import fitz

from docx import Document

from ocr import extract_text_from_pil_image


def get_extension(filename):
    if "." not in filename:
        return ""

    return filename.rsplit(
        ".",
        1
    )[1].lower()


def extract_from_txt(file_path):
    try:
        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:
            text = file.read()

        return {
            "success": True,
            "text": text,
            "method": "TXT extraction"
        }

    except Exception as error:
        return {
            "success": False,
            "error": str(error)
        }


def extract_from_docx(file_path):
    try:
        document = Document(
            file_path
        )

        sections = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                sections.append(text)

        for table_index, table in enumerate(
            document.tables,
            start=1
        ):
            sections.append(
                f"\n[TABLE {table_index}]"
            )

            for row in table.rows:
                cells = []

                for cell in row.cells:
                    cells.append(
                        cell.text.strip()
                    )

                sections.append(
                    " | ".join(cells)
                )

        return {
            "success": True,
            "text": "\n".join(sections),
            "method": "DOCX extraction"
        }

    except Exception as error:
        return {
            "success": False,
            "error": str(error)
        }


def extract_from_pdf(file_path):
    try:
        document = fitz.open(
            file_path
        )

        pages = []
        total_ocr_pages = 0

        for page_number, page in enumerate(
            document,
            start=1
        ):
            direct_text = page.get_text(
                "text"
            ).strip()

            if len(direct_text) >= 30:
                pages.append(
                    f"[PAGE {page_number}]\n"
                    f"{direct_text}"
                )
                continue

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(
                    2,
                    2
                ),
                alpha=False
            )

            image_bytes = pixmap.tobytes(
                "png"
            )

            from PIL import Image
            from io import BytesIO

            image = Image.open(
                BytesIO(image_bytes)
            )

            ocr_result = extract_text_from_pil_image(
                image
            )

            if ocr_result["success"]:
                total_ocr_pages += 1

                pages.append(
                    f"[PAGE {page_number}]\n"
                    f"{ocr_result['text']}"
                )
            else:
                pages.append(
                    f"[PAGE {page_number}]\n"
                    f"[OCR ERROR: "
                    f"{ocr_result['error']}]"
                )

        document.close()

        method = "PDF text extraction"

        if total_ocr_pages > 0:
            method += (
                f" + OCR "
                f"({total_ocr_pages} scanned page(s))"
            )

        return {
            "success": True,
            "text": "\n\n".join(pages),
            "method": method
        }

    except Exception as error:
        return {
            "success": False,
            "error": str(error)
        }


def process_document(
    file_path
):
    extension = get_extension(
        os.path.basename(file_path)
    )

    image_extensions = {
        "jpg",
        "jpeg",
        "png",
        "webp"
    }

    if extension in image_extensions:
        from ocr import extract_text_from_image

        return extract_text_from_image(
            file_path
        )

    if extension == "pdf":
        return extract_from_pdf(
            file_path
        )

    if extension == "docx":
        return extract_from_docx(
            file_path
        )

    if extension == "txt":
        return extract_from_txt(
            file_path
        )

    return {
        "success": False,
        "error": (
            f"Unsupported file type: "
            f".{extension}"
        )
    }