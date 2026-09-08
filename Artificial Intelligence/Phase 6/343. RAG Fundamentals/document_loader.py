from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
    ".csv"
}


def load_txt(path):
    for encoding in (
        "utf-8",
        "utf-8-sig",
        "latin-1"
    ):
        try:
            return Path(path).read_text(
                encoding=encoding
            )
        except UnicodeDecodeError:
            continue

    raise ValueError(
        "Unable to decode TXT file."
    )


def load_pdf(path):
    reader = PdfReader(path)
    pages = []

    for number, page in enumerate(
        reader.pages,
        start=1
    ):
        text = page.extract_text()

        if text:
            pages.append(
                f"[Page {number}]\n{text}"
            )

    return "\n\n".join(pages)


def load_docx(path):
    document = Document(path)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for number, table in enumerate(
        document.tables,
        start=1
    ):
        parts.append(
            f"[Table {number}]"
        )

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            parts.append(
                " | ".join(values)
            )

    return "\n".join(parts)


def load_csv(path):
    dataframe = pd.read_csv(path)
    dataframe = dataframe.fillna("")
    parts = []

    for _, row in dataframe.iterrows():
        values = []

        for column in dataframe.columns:
            value = str(
                row[column]
            ).strip()

            if value:
                values.append(
                    f"{column}: {value}"
                )

        if values:
            parts.append(
                " | ".join(values)
            )

    return "\n".join(parts)


def clean_text(text):
    lines = []

    for line in text.splitlines():
        line = " ".join(
            line.split()
        )

        if line:
            lines.append(line)

    return "\n".join(lines)


def load_document(path):
    path = Path(path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".txt":
        text = load_txt(path)

    elif extension == ".pdf":
        text = load_pdf(path)

    elif extension == ".docx":
        text = load_docx(path)

    elif extension == ".csv":
        text = load_csv(path)

    else:
        raise ValueError(
            "Unsupported document."
        )

    return clean_text(text)